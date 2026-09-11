#!/usr/bin/env python3
"""Replace the manuscript's nine figures with NDSS-style monochrome assets."""

from __future__ import annotations

import os
import struct
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "qkl" / "docs" / "xn" / "论文初稿-中文版_仓库版.docx"
BACKUP = ROOT / "qkl" / "docs" / "xn" / "论文初稿-中文版_仓库版_图形优化前备份.docx"
FIGURES = ROOT / "qkl" / "docs" / "xn" / "figures_ndss" / "final"

# DOCX media part -> exact replacement asset.
PART_TO_FIGURE = {
    "image1.png": "fig1_architecture_bw.png",
    "image3.png": "fig2_lifecycle_bw.png",
    "image4.png": "fig3_byzantine_bw.png",
    "image10.png": "fig4_correlated_failure_bw.png",
    "image5.png": "figA1_simulation_bw.png",
    "image6.png": "figA2_ablation_bw.png",
    "image7.png": "figA3_calibration_bw.png",
    "image8.png": "figA4_decentralization_bw.png",
    "image9.png": "figA5_onchain_cost_bw.png",
}


def png_dimensions(path: Path) -> tuple[int, int]:
    with path.open("rb") as stream:
        header = stream.read(24)
    if header[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"Not a PNG: {path}")
    return struct.unpack(">II", header[16:24])


def main() -> None:
    if not DOCX.exists():
        raise FileNotFoundError(DOCX)
    missing = [FIGURES / name for name in PART_TO_FIGURE.values() if not (FIGURES / name).exists()]
    if missing:
        raise FileNotFoundError(f"Missing figure assets: {missing}")

    # Keep the first pre-figure-edit snapshot immutable.
    if not BACKUP.exists():
        BACKUP.write_bytes(DOCX.read_bytes())

    document = Document(DOCX)
    replaced: set[str] = set()
    display_width_inches = 7.05

    for shape in document.inline_shapes:
        blips = shape._inline.xpath(".//a:blip")
        if not blips:
            continue
        relationship_id = blips[0].get(qn("r:embed"))
        image_part = document.part.related_parts[relationship_id]
        part_name = Path(str(image_part.partname)).name
        replacement_name = PART_TO_FIGURE.get(part_name)
        if replacement_name is None:
            continue

        replacement = FIGURES / replacement_name
        pixel_width, pixel_height = png_dimensions(replacement)
        image_part._blob = replacement.read_bytes()
        shape.width = Inches(display_width_inches)
        shape.height = Inches(display_width_inches * pixel_height / pixel_width)
        replaced.add(part_name)

    expected = set(PART_TO_FIGURE)
    if replaced != expected:
        raise RuntimeError(f"Figure mapping mismatch: replaced={sorted(replaced)}, expected={sorted(expected)}")

    # The manuscript's document default is an exact 11-pt line height.  Without
    # an override, Word clips inline pictures to that text-line height.  Apply an
    # automatic line height only to the nine picture paragraphs and keep each
    # picture with its following caption.
    picture_paragraphs = 0
    for paragraph in document.paragraphs:
        if not paragraph._p.xpath(".//w:drawing"):
            continue
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.keep_with_next = True
        picture_paragraphs += 1
    if picture_paragraphs != 9:
        raise RuntimeError(f"Expected 9 picture paragraphs, found {picture_paragraphs}")

    temporary = DOCX.with_name(f".{DOCX.stem}.ndss-figures.tmp.docx")
    document.save(temporary)

    # Reopen before replacing the requested manuscript to catch package errors.
    check = Document(temporary)
    if len(check.inline_shapes) != 9:
        temporary.unlink(missing_ok=True)
        raise RuntimeError(f"Expected 9 inline figures after save, found {len(check.inline_shapes)}")

    os.replace(temporary, DOCX)
    print(f"Updated {DOCX}")
    print(f"Backup  {BACKUP}")
    print("Replaced 9 figures, preserved aspect ratios, and enabled automatic picture-line height")


if __name__ == "__main__":
    main()
