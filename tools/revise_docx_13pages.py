#!/usr/bin/env python3
"""Restructure the Chinese manuscript to meet a 13-page main-text limit.

The script operates on the pre-edit backup and writes a candidate DOCX.  It
preserves the original XML nodes for moved paragraphs, tables, drawings, and
section breaks so that Word's two-column/full-width layout remains intact.
"""

from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph


def normalized_text(element) -> str:
    if not isinstance(element, CT_P):
        return ""
    return " ".join(Paragraph(element, element.getroottree()).text.split())


def paragraph_text(element) -> str:
    if not isinstance(element, CT_P):
        return ""
    return "".join(element.xpath(".//w:t/text()"))


def replace_in_runs(element, replacements: list[tuple[str, str]]) -> None:
    """Replace paragraph text while retaining the surrounding run formatting.

    Cross-run matches are assigned to the first affected run; unaffected text
    remains in its original runs.
    """
    if not isinstance(element, CT_P):
        return
    for old, new in replacements:
        while True:
            text_nodes = element.xpath(".//w:t")
            values = [node.text or "" for node in text_nodes]
            combined = "".join(values)
            start = combined.find(old)
            if start < 0:
                break
            end = start + len(old)

            offsets = []
            cursor = 0
            for value in values:
                offsets.append(cursor)
                cursor += len(value)

            start_index = next(
                i for i, (offset, value) in enumerate(zip(offsets, values))
                if offset <= start < offset + len(value)
            )
            end_index = next(
                i for i, (offset, value) in enumerate(zip(offsets, values))
                if offset <= end - 1 < offset + len(value)
            )
            start_offset = start - offsets[start_index]
            end_offset = end - offsets[end_index]

            if start_index == end_index:
                text_nodes[start_index].text = (
                    values[start_index][:start_offset]
                    + new
                    + values[start_index][end_offset:]
                )
            else:
                text_nodes[start_index].text = values[start_index][:start_offset] + new
                for i in range(start_index + 1, end_index):
                    text_nodes[i].text = ""
                text_nodes[end_index].text = values[end_index][end_offset:]


def set_paragraph_text(element, text: str) -> None:
    """Replace paragraph content but retain paragraph and first-run styling."""
    if not isinstance(element, CT_P):
        raise TypeError("Expected a paragraph element")

    first_rpr = None
    first_run = element.find(qn("w:r"))
    if first_run is not None:
        rpr = first_run.find(qn("w:rPr"))
        if rpr is not None:
            first_rpr = copy.deepcopy(rpr)

    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)

    run = OxmlElement("w:r")
    if first_rpr is not None:
        run.append(first_rpr)
    text_node = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        text_node.set(qn("xml:space"), "preserve")
    text_node.text = text
    run.append(text_node)
    element.append(run)


def clone_heading(source, text: str, page_break_before: bool = False):
    heading = copy.deepcopy(source)
    set_paragraph_text(heading, text)
    if page_break_before:
        ppr = heading.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            heading.insert(0, ppr)
        if ppr.find(qn("w:pageBreakBefore")) is None:
            ppr.append(OxmlElement("w:pageBreakBefore"))
    return heading


def locate_unique(blocks, exact_text: str):
    hits = [b for b in blocks if paragraph_text(b).strip() == exact_text]
    if len(hits) != 1:
        raise RuntimeError(f"Expected one paragraph {exact_text!r}; found {len(hits)}")
    return hits[0]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    doc = Document(args.source)
    body = doc.element.body
    blocks = list(body.iterchildren())
    paragraphs = [p._p for p in doc.paragraphs]
    if len(blocks) != 565:
        raise RuntimeError(f"Unexpected document structure: {len(blocks)} body blocks")

    # Anchor validation protects against applying fixed structural ranges to a
    # different manuscript revision.
    expected = {
        29: "表 A 从两条研究脉络出发",
        49: "",
        385: "D.2 概率性语义层下的容错(仿真)",
        397: "D.3 消融:优势来源",
        416: "D.4 策略性对手与串通检测",
        475: "D.11 链上成本(FISCO-BCOS 实测)",
        492: "D.12 模型异构性与相关失败(Q1 补充实验)",
        504: "D.13 真实判断驱动的语义层(D.2 的真实数据对照)",
        513: "D.14 伪造拒绝攻击(对称对照,场景 2)",
        520: "D.15 分阶段延迟分解(off-chain 实测 + 链上锚定实测)",
        527: "D.16 小结",
        541: "IX. 结论",
        545: "参考文献",
    }
    for index, prefix in expected.items():
        text = paragraph_text(blocks[index]).strip()
        if prefix and not text.startswith(prefix):
            raise RuntimeError(f"Block {index} mismatch: {text!r}")

    # Elements that will form appendix subsections.  Ranges are half-open and
    # were mapped against the immutable pre-edit backup.
    appendix_groups = [
        ("A.1 单因子加权基线补充结果", blocks[376:381]),
        ("A.2 概率性语义层下的容错（仿真）", blocks[385:397]),
        ("A.3 扩展消融结果", blocks[407:416]),
        ("A.4 策略性对手与串通检测", blocks[416:428]),
        ("A.5 吞吐与可扩展性", blocks[428:434]),
        ("A.6 沉默攻击", blocks[434:440]),
        ("A.7 参数敏感性", blocks[440:446]),
        ("A.8 激励相容的经验验证", blocks[446:452]),
        ("A.9 置信度校准", blocks[452:464]),
        ("A.10 去中心化与集中化风险", blocks[464:475]),
        ("A.11 链上成本扩展对照", blocks[484:492]),
        ("A.12 模型相关失败的完整数值", blocks[495:499]),
        ("A.13 伪造拒绝攻击", blocks[513:520]),
        ("A.14 分阶段延迟分解", blocks[520:527]),
    ]

    moved = {id(e) for _, group in appendix_groups for e in group}
    if sum(len(group) for _, group in appendix_groups) != len(moved):
        raise RuntimeError("Appendix block ranges overlap")

    # Delete the redundant related-work matrix, the redundant sequence figure,
    # the merged D.13 heading, and the results recap that duplicated conclusions.
    delete_ranges = [blocks[29:35], blocks[49:53], blocks[504:505], blocks[527:541]]
    deleted = {id(e) for group in delete_ranges for e in group}
    if moved & deleted:
        raise RuntimeError("A block cannot be both moved and deleted")

    for element in list(blocks):
        if id(element) in moved or id(element) in deleted:
            body.remove(element)

    # Main-text section headings after consolidation.
    set_paragraph_text(blocks[397], "D.2 核心机制消融")
    set_paragraph_text(blocks[475], "D.3 链上成本")
    set_paragraph_text(blocks[492], "D.4 真实模型异构性与语义判断")

    # Rewrite only the paragraphs whose function changed; numerical claims and
    # experimental conditions remain unchanged.
    replace_in_runs(
        paragraphs[4],
        [
            (
                "基于真实智能体的实验显示，在 40% 恶意智能体下",
                "经真实 LLM 校准的概率语义层仿真显示，在 40% 恶意智能体下",
            ),
            (
                "；当恶意智能体占调用委员会多数（60%）时",
                "；真实智能体对抗实验进一步表明，当恶意智能体占调用委员会多数（60%）时",
            ),
        ],
    )
    set_paragraph_text(paragraphs[43], "系统整体架构如图 1 所示。")

    targeted_replacements = {
        139: [("第 VIII-D.3 节", "第 VIII-D.2 节")],
        210: [("第 VIII-D.4 节的", "附录 A.4 的")],
        265: [("第 VIII-D.4 节", "附录 A.4")],
        311: [
            (
                "第 VIII-D.2/D.3 节在概率语义层实测的 false-accept（本方案较各基线低约一个数量级）"
                "与式(7)(8)预测的指数压制相符；推论2的乘性增益对应 D.3 消融中“关闭仲裁后正确率"
                "由 95.9% 跌至 15.9%”的结果。",
                "附录 A.2 的概率语义层结果和第 VIII-D.2 节的机制消融均显示 false-accept 受到约一个"
                "数量级的压制，与式(7)(8)的预测一致；推论2的乘性增益对应第 VIII-D.2 节中关闭仲裁"
                "后正确率由 95.9% 降至 15.9% 的结果。",
            ),
        ],
        326: [("第 VIII-D.5 节实测", "附录 A.5 的实测")],
        347: [
            ("D.2 即在此值下报告", "附录 A.2 在此值下报告"),
            ("D.7 进一步", "附录 A.7 进一步"),
            ("第 VIII-D.12/D.13 节", "第 VIII-D.4 节"),
            ("D.13 进一步", "第 VIII-D.4 节进一步"),
        ],
        362: [("见 D.10", "见附录 A.10"), ("见 D.9", "见附录 A.9")],
        381: [("第 VIII-D.13 节", "第 VIII-D.4 节")],
        394: [("第 VIII-D.12 节", "第 VIII-D.4 节")],
        486: [
            ("D.2 用", "附录 A.2 采用"),
            ("第 VIII-D.12 节", "本节前述异构性实验"),
        ],
        491: [
            ("见 D.12", "见本节前述结果"),
            ("显著高于 D.2 在", "显著高于附录 A.2 在"),
            ("也说明 D.2 的", "也说明附录 A.2 给出的"),
        ],
        500: [("D.11 实测", "第 VIII-D.3 节实测"), ("见 D.11", "见第 VIII-D.3 节")],
        502: [("链上锚定为 D.11 实测", "链上锚定采用第 VIII-D.3 节实测值")],
    }
    for index, replacements in targeted_replacements.items():
        replace_in_runs(paragraphs[index], replacements)

    set_paragraph_text(
        paragraphs[363],
        "其中 B3/B6/B4 三个加权基线与本方案消费同一批真实判断，用于隔离信誉加权、"
        "置信度加权以及 rep×conf 联合加权与仲裁升级的贡献；B3/B6 的完整结果见附录 A.1，"
        "B4 的对照见第 VIII-D.1 节与附录 A.2。",
    )
    set_paragraph_text(
        paragraphs[365],
        "所有正确率和误受率均为实测值；配对比较中，各方法消费同一批智能体判断，显著性采用"
        "配对 bootstrap 95% 置信区间与符号检验。正文聚焦拜占庭鲁棒性、核心机制消融、链上成本"
        "以及真实模型异构性与语义判断；策略性与沉默攻击、参数敏感性、置信度校准、去中心化、"
        "吞吐及分阶段延迟等补充结果见附录 A。",
    )

    # Rename headings of moved full subsections.  Partial groups receive cloned
    # headings below rather than reusing captions as structural labels.
    moved_heading_names = {
        385: "A.2 概率性语义层下的容错（仿真）",
        416: "A.4 策略性对手与串通检测",
        428: "A.5 吞吐与可扩展性",
        434: "A.6 沉默攻击",
        440: "A.7 参数敏感性",
        446: "A.8 激励相容的经验验证",
        452: "A.9 置信度校准",
        464: "A.10 去中心化与集中化风险",
        513: "A.13 伪造拒绝攻击",
        520: "A.14 分阶段延迟分解",
    }
    for index, text in moved_heading_names.items():
        set_paragraph_text(blocks[index], text)

    # Renumber every retained/moved display through collision-free placeholders.
    table_map = {
        "表 2b": "表 A1",
        "表 3": "表 A2",
        "表 4": "表 3",
        "表 B": "表 A3",
        "表 5": "表 A4",
        "表 6": "表 A5",
        "表 7": "表 A6",
        "表 8": "表 A7",
        "表 9": "表 A8",
        "表 10": "表 A9",
        "表 11": "表 A10",
        "表 12": "表 A11",
        "表 13": "表 4",
        "表 14": "表 A12",
        "表 15": "表 A13",
        "表 16": "表 5",
        "表 17": "表 A14",
        "表 18": "表 A15",
    }
    figure_map = {
        "图 3": "图 2",
        "图 4": "图 3",
        "图 5": "图 A1",
        "图 6": "图 A2",
        "图 7": "图 A3",
        "图 8": "图 A4",
        "图 9": "图 A5",
        "图 10": "图 4",
    }
    surviving = [e for e in blocks if id(e) not in deleted]
    for n, (old, new) in enumerate(list(table_map.items()) + list(figure_map.items())):
        token = f"@@DISPLAY_{n}@@"
        for element in surviving:
            replace_in_runs(element, [(old, token)])
        for element in surviving:
            replace_in_runs(element, [(token, new)])

    # The ablation and cost prose now explicitly signal that their large visual
    # duplicates are in the appendix.
    replace_in_runs(paragraphs[393], [("(表 3;图 A2)", "（表 3；图示见附录图 A2）")])
    replace_in_runs(paragraphs[466], [("对照(表 A12、图 A5)", "扩展对照见附录表 A12 与图 A5")])
    replace_in_runs(paragraphs[476], [("结果见表 A13 与图 4", "结果见图 4，完整数值见附录表 A13")])

    # Keep references immediately after the conclusion (the end of the main
    # text), then place the appendix after the references.  The appendix page
    # break makes the main-text/reference/appendix boundary unambiguous.
    reference = locate_unique(list(body.iterchildren()), "参考文献")
    terminal_section = body.sectPr
    main_heading_template = blocks[541]
    subheading_template = blocks[397]
    terminal_section.addprevious(clone_heading(main_heading_template, "附录A 补充实验", True))

    full_heading_elements = {blocks[i] for i in moved_heading_names}
    for title, group in appendix_groups:
        if group and group[0] in full_heading_elements:
            # The original heading was renamed above and moves with the group.
            pass
        else:
            terminal_section.addprevious(clone_heading(subheading_template, title))
        for element in group:
            terminal_section.addprevious(element)

    # Final textual integrity checks before serialization.
    final_text = "\n".join(paragraph_text(e) for e in body.iterchildren() if isinstance(e, CT_P))
    forbidden = [
        "D.16 小结",
        "表 A. 相关工作横向对比",
        "图 2. 智能合约到链下智能体的调用时序",
        "@@DISPLAY_",
    ]
    for marker in forbidden:
        if marker in final_text:
            raise RuntimeError(f"Stale marker remains: {marker}")
    required = [
        "D.1 真智能体拜占庭鲁棒性",
        "D.2 核心机制消融",
        "D.3 链上成本",
        "D.4 真实模型异构性与语义判断",
        "附录A 补充实验",
        "A.14 分阶段延迟分解",
        "参考文献",
    ]
    for marker in required:
        if marker not in final_text:
            raise RuntimeError(f"Required marker missing: {marker}")
    if not (
        final_text.index("IX. 结论")
        < final_text.index("参考文献")
        < final_text.index("附录A 补充实验")
    ):
        raise RuntimeError("Expected order is conclusion, references, then appendix")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
