#!/usr/bin/env python3
"""Build the clean revised manuscript from the annotated DOCX and experiment JSON."""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "docs" / "xn"
SOURCE = DOC_DIR / "Electronics_异构联盟链可审计语义查询框架_逐段起稿.docx"
RESULTS = DOC_DIR / "experiments" / "semantic_query_framework_results.json"
LIVE_ROOT = DOC_DIR / "experiments" / "semantic-query-live-runs" / "2026-08-25T01-46-44-989Z"
PHYSICAL_ROOT = DOC_DIR / "experiments" / "semantic-query-physical-fault-runs" / "2026-08-25T01-52-36-106Z"
EXTENDED_RESULTS = DOC_DIR / "experiments" / "semantic-query-extended-evaluation" / "results.json"
HISTORICAL_AUDIT = DOC_DIR / "experiments" / "independent_audit_verification.json"
V2_AUDIT = DOC_DIR / "experiments" / "semantic-query-receipt-v2-independent-verification.json"
SATURATION_ROOT = DOC_DIR / "experiments" / "semantic-query-live-concurrency-runs" / "2026-09-01T01-28-00-905Z"
REGULATORY_ROOT = DOC_DIR / "experiments" / "semantic-query-live-regulatory-runs" / "2026-09-01T01-38-58-803Z"
RESILIENCE_ROOT = DOC_DIR / "experiments" / "semantic-query-live-resilience-runs" / "2026-09-01T07-55-45-402Z"
MATERIAL_OUTAGE_ROOT = DOC_DIR / "experiments" / "semantic-query-live-resilience-runs" / "2026-09-02T00-46-44-501Z"
MULTIWALLET_ROOT = DOC_DIR / "experiments" / "semantic-query-live-multiwallet-runs" / "2026-09-02-formal-w4"
MODEL_BENCHMARK_ROOT = DOC_DIR / "experiments" / "semantic-model-benchmark-120"
OUTPUT = DOC_DIR / "Electronics_异构联盟链可审计语义查询框架_统一实证修订版.docx"
CHANGELOG = DOC_DIR / "Electronics_异构联盟链可审计语义查询框架_统一实证修订说明.md"
FIGURE_DIR = DOC_DIR / "figures"


def find_paragraph(doc: Document, prefix: str):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def set_text(paragraph, text: str, *, bold_prefix: str | None = None):
    paragraph.clear()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true")
    tr_pr.append(tbl_header)


def insert_caption_and_table(doc: Document, anchor, caption: str, headers, rows, widths=None):
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_p.add_run(caption)
    run.bold = True
    run.font.size = Pt(9)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = str(value)
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.size = Pt(8)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = str(value)
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = width

    anchor._p.addnext(caption_p._p)
    caption_p._p.addnext(table._tbl)
    return table


def insert_picture(doc: Document, anchor, image_path: Path, caption: str, width_inches: float = 6.6):
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.add_run().add_picture(str(image_path), width=Inches(width_inches))
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_p.add_run(caption)
    run.bold = True
    run.font.size = Pt(9)
    anchor._p.addnext(image_p._p)
    image_p._p.addnext(caption_p._p)
    return image_p


def anchor_statistics(database_path: Path):
    connection = sqlite3.connect(database_path)
    try:
        rows = connection.execute("SELECT block_height, anchor_json FROM audit_anchors").fetchall()
    finally:
        connection.close()
    records = [json.loads(row[1]) for row in rows]
    gas = [int(record["gasUsed"]) for record in records if record.get("gasUsed") is not None]
    return {
        "count": len(rows),
        "gasCount": len(gas),
        "gasMean": sum(gas) / len(gas) if gas else None,
        "gasMin": min(gas) if gas else None,
        "gasMax": max(gas) if gas else None,
        "heightMin": min(row[0] for row in rows),
        "heightMax": max(row[0] for row in rows),
    }


def pct(value: float, digits=1):
    return f"{value * 100:.{digits}f}%"


def main():
    report = json.loads(RESULTS.read_text(encoding="utf-8"))
    live = json.loads((LIVE_ROOT / "summary.json").read_text(encoding="utf-8"))
    physical = json.loads((PHYSICAL_ROOT / "summary.json").read_text(encoding="utf-8"))
    extended = json.loads(EXTENDED_RESULTS.read_text(encoding="utf-8"))
    historical_audit = json.loads(HISTORICAL_AUDIT.read_text(encoding="utf-8"))
    v2_audit = json.loads(V2_AUDIT.read_text(encoding="utf-8"))
    saturation = json.loads((SATURATION_ROOT / "summary.json").read_text(encoding="utf-8"))
    saturation_audit = json.loads((SATURATION_ROOT / "independent_verification.json").read_text(encoding="utf-8"))
    regulatory = json.loads((REGULATORY_ROOT / "summary.json").read_text(encoding="utf-8"))
    regulatory_audit = json.loads((REGULATORY_ROOT / "independent_verification.json").read_text(encoding="utf-8"))
    resilience = json.loads((RESILIENCE_ROOT / "summary.json").read_text(encoding="utf-8"))
    resilience_audit = json.loads((RESILIENCE_ROOT / "independent_verification.json").read_text(encoding="utf-8"))
    material_outage = json.loads((MATERIAL_OUTAGE_ROOT / "summary.json").read_text(encoding="utf-8"))
    material_outage_audit = json.loads((MATERIAL_OUTAGE_ROOT / "independent_verification.json").read_text(encoding="utf-8"))
    multiwallet = json.loads((MULTIWALLET_ROOT / "summary.json").read_text(encoding="utf-8"))
    multiwallet_comparison = json.loads((MULTIWALLET_ROOT / "comparison.json").read_text(encoding="utf-8"))
    multiwallet_audit = json.loads((MULTIWALLET_ROOT / "independent_verification.json").read_text(encoding="utf-8"))
    model_cases = json.loads((MODEL_BENCHMARK_ROOT / "cases.json").read_text(encoding="utf-8"))
    live_anchor = anchor_statistics(LIVE_ROOT / "audit.db")
    physical_anchor = anchor_statistics(PHYSICAL_ROOT / "audit.db")
    saturation_anchors = sum(anchor_statistics(SATURATION_ROOT / f"c{level}" / "audit.db")["count"] for level in [1, 2, 4, 8, 16])
    regulatory_anchor = anchor_statistics(REGULATORY_ROOT / "audit.db")
    resilience_anchor = anchor_statistics(RESILIENCE_ROOT / "audit.db")
    material_outage_anchor = anchor_statistics(MATERIAL_OUTAGE_ROOT / "audit.db")
    multiwallet_anchor_stats = [
        anchor_statistics(MULTIWALLET_ROOT / f"c{level}" / "audit.db")
        for level in [1, 2, 4, 8, 16]
    ]
    multiwallet_anchors = sum(item["count"] for item in multiwallet_anchor_stats)
    doc = Document(SOURCE)

    # Remove all single-cell author-note/evidence-gap boxes while preserving the
    # substantive unified-evidence-field table.
    removed_notes = 0
    for table in list(doc.tables):
        if len(table.rows) == 1 and len(table.columns) == 1:
            table._element.getparent().remove(table._element)
            removed_notes += 1

    rq1 = report["rq1EvidenceValidation"]
    rq2_rows = report["rq2DagScheduling"]["rows"]
    rq3 = report["rq3FaultRecovery"]
    rq4 = report["rq4AuditReconstruction"]
    pruning = report["rq5PruningAblation"]
    chain = report["archivedOnchain"]
    meta = report["meta"]
    live_by = live["byScenario"]
    physical_by = physical["byScenario"]
    historical_live, historical_physical = historical_audit["results"]
    v2_result = v2_audit["results"][0]
    saturation_audit_results = saturation_audit["results"]
    regulatory_audit_result = regulatory_audit["results"][0]
    resilience_audit_result = resilience_audit["results"][0]
    material_outage_audit_result = material_outage_audit["results"][0]
    multiwallet_audit_results = multiwallet_audit["results"]
    live_v2_queries = (
        sum(row["queries"] for row in saturation_audit_results)
        + regulatory_audit_result["queries"]
        + resilience_audit_result["queries"]
        + material_outage_audit_result["queries"]
        + sum(row["queries"] for row in multiwallet_audit_results)
    )
    live_v2_evidence = sum(
        row["bindingCoverage"]["successfulEvidenceNodes"] for row in saturation_audit_results
    ) + regulatory_audit_result["bindingCoverage"]["successfulEvidenceNodes"] + resilience_audit_result["bindingCoverage"]["successfulEvidenceNodes"] + material_outage_audit_result["bindingCoverage"]["successfulEvidenceNodes"] + sum(row["bindingCoverage"]["successfulEvidenceNodes"] for row in multiwallet_audit_results)
    live_v2_services = sum(
        row["bindingCoverage"]["successfulSemanticNodes"] for row in saturation_audit_results
    ) + regulatory_audit_result["bindingCoverage"]["successfulSemanticNodes"] + resilience_audit_result["bindingCoverage"]["successfulSemanticNodes"] + material_outage_audit_result["bindingCoverage"]["successfulSemanticNodes"] + sum(row["bindingCoverage"]["successfulSemanticNodes"] for row in multiwallet_audit_results)
    live_v2_mutations = live_v2_queries * 4
    total_formal_anchors = (
        live_anchor["count"]
        + physical_anchor["count"]
        + saturation_anchors
        + regulatory_anchor["count"]
        + resilience_anchor["count"]
        + material_outage_anchor["count"]
        + multiwallet_anchors
    )
    resilience_by = resilience["byScenario"]
    material_outage_result = material_outage["byScenario"]["material_store_outage"]
    saturation_throughputs = ", ".join(f"{row['throughput']:.3f}" for row in saturation["rows"])
    multiwallet_throughputs = ", ".join(f"{row['throughput']:.3f}" for row in multiwallet["rows"])
    multiwallet_by_concurrency = {
        row["concurrency"]: row for row in multiwallet_comparison["rows"]
    }
    historical_binding_rate = (
        historical_live["bindingCoverage"]["coveredEvidenceNodes"]
        + historical_physical["bindingCoverage"]["coveredEvidenceNodes"]
    ) / (
        historical_live["bindingCoverage"]["successfulEvidenceNodes"]
        + historical_physical["bindingCoverage"]["successfulEvidenceNodes"]
    )

    # Abstract, contributions, and roadmap.
    abstract = (
        "摘  要：异构联盟链能够传递远端交易或状态，但供应链融资、监管审计等业务需要组合多条链及链下材料，判断带有主体、金额、时序和有效期约束的业务谓词。现有机制缺少对多源证据的统一表示、依赖编排、失败处置和结果追溯，直接让链下语义服务产生合约可消费结果会引入证据错配和输出不可重构风险。为此，本文提出面向异构联盟链的可审计跨链语义查询框架，将Fabric、FISCO-BCOS及链下材料规范化为与查询和业务谓词绑定的证据对象，将查询编译为证据依赖图，并通过并行取证、确定性预验证、受证据约束的语义服务和有界恢复形成三值业务判断。完整证据包与执行轨迹保存在链下，链上仅锚定回执摘要。在单机双链环境中，600个正常与受控故障查询均进入预期终态，正常查询p95为580.407 ms；40次容器级节点暂停均完成候补端点切换。网络分区、FISCO恢复追赶、相关双节点暂停、材料存储不可用和审计持久化故障各5次，25/25次得到预期恢复或保守终态。单钱包真实并发实验完成300次查询，在并发1–16下吞吐为1.854–1.886 query/s并于并发2达到饱和判据；四钱包对照另完成300次查询，吞吐在并发4达到7.150 query/s，饱和点移至并发8。监管披露第二负载的150次真实链查询全部符合预期终态。正式数据集共含1,454笔专用锚定交易，gasUsed均为35,166。独立Python审计器对历史643份回执和修复格式811份真实链回执全部完成重构、锚点匹配及四类篡改检测；修复格式的4,010个成功证据节点和776个服务输出均为显式绑定。结果验证了可审计执行、所测故障恢复、锚定并发优化和格式修复效果，但不代表跨地域性能上限、一般拜占庭容错或语义真值保证。"
    )
    set_text(find_paragraph(doc, "摘  要："), abstract, bold_prefix="摘  要：")
    set_text(
        find_paragraph(doc, "（4）异构原型与面向机制的评估方案。"),
        "（4）异构原型与分层实证评估。在Fabric和FISCO-BCOS之间实现真实取证、DAG执行、SQLite审计持久化和FISCO回执根锚定，并通过第二业务负载、受控故障、容器级节点暂停、网络分区、相关节点故障、恢复追赶、材料与审计存储故障、独立回执审计、DAG扩展和机制消融分别评估功能、恢复、审计与开销，同时明确区分真实双链实测、本地控制实验和探索性仿真。",
        bold_prefix="（4）异构原型与分层实证评估。",
    )
    set_text(
        find_paragraph(doc, "本文其余部分组织如下："),
        "本文其余部分组织如下：第2节讨论跨链互操作、预言机、业务流程和数据溯源研究；第3节给出系统模型与问题定义；第4节介绍框架机制；第5节分析关键性质和复杂度；第6节描述原型实现；第7节报告功能、调度、故障恢复、审计和消融实验；第8节讨论适用边界和有效性威胁；第9节总结全文。",
    )

    # Prototype implementation: separate the integrated two-chain path from the
    # newly executed off-chain reference harness.
    set_text(
        find_paragraph(doc, "我们实现了一个连接Hyperledger Fabric与FISCO-BCOS的原型"),
        "我们实现了一个连接Hyperledger Fabric与FISCO-BCOS的跨链原型。供应链融资查询从Fabric Gateway和FISCO合约读取四类账本证据，并与链下质量文档共同进入10节点DAG；监管披露查询则把实体、许可和披露映射到Fabric状态，把交易和政策映射到FISCO合约状态，使五个证据叶均来自真实链。结果生成稳定序列化审计回执，持久化到SQLite，并通过SemanticQueryAuditAnchor合约将回执根提交至FISCO-BCOS。因而，两类业务中的证据适配、DAG调度、异常终态、审计持久化和链上锚定均在真实双链路径中连通。当前语义节点仍使用确定性本地服务，并且Fabric定位记录是查询时账本高度与业务资源，而非交易包含证明。",
    )
    set_text(
        find_paragraph(doc, "链下执行层以Node.js服务组织"),
        "链下执行层以Node.js服务组织，Fabric和FISCO适配器均支持候补端点。查询调度器按拓扑依赖执行证据、确定性规则、语义与根节点，并为每个节点维护候选集、尝试次数、超时和全局deadline。锚定层支持单钱包顺序提交，也支持将查询按轮询分配至多个签名钱包；每个钱包内部仍串行化交易，以避免同一账户的nonce竞争。实验驱动另提供顺序、无依赖全并行及依赖感知三种微基准模式。节点重试、候补切换、证据不足与执行失败均记入可哈希轨迹。",
    )
    set_text(
        find_paragraph(doc, "确定性验证以纯函数或版本化规则模块实现"),
        "确定性验证以版本化函数实现，覆盖必需字段、查询标识、源链、状态位置、终局性、时效、模式、负载摘要、证据标识和Ed25519签名。实验中的语义节点使用受控本地服务时间，以隔离调度开销；因此其结果只用于比较编排策略，不代表真实LLM或远端RPC时延。现有语义服务接口和双模型记录仅用于探索模型相关错误，不作为端到端性能数据。",
    )
    set_text(
        find_paragraph(doc, "每次查询生成一份JSON审计回执"),
        "每次实验查询生成一份稳定序列化的JSON审计回执，记录查询定义、DAG版本、节点依赖、输入引用、输出摘要、处理版本、状态、时间和Ed25519签名。回执整体计算SHA-256根；审计器从根节点逆向检查图闭包、输入摘要、签名和链上锚定根，并对缺失记录返回明确断点。独立复核发现历史643份链上锚定回执虽能验证整体根，但本地文档证据未全部进入evidence表，证据及服务记录也缺少显式nodeId。修复后的v2格式为全部证据和服务输出增加节点绑定；历史锚点保持原样，不作追溯改写。敏感业务负载仍应在实际部署中与摘要分离并受访问控制。",
    )
    set_text(
        find_paragraph(doc, "为支持复现实验，测试驱动器应固定"),
        f"离线实验固定随机种子{meta['baseSeed']}；第一轮真实双链基准在3次预热后对6个场景各执行5轮、每轮20个查询。物理故障实验对Fabric Org1 peer和FISCO node0各执行20次暂停—恢复。单钱包与四钱包并发实验均在并发1/2/4/8/16下各执行3轮、每轮20个计量查询和2个预热查询，并按1 s间隔采集容器资源。监管披露真实链实验对5个场景各执行3轮、每轮10个计量查询。韧性实验对Fabric网络分区、FISCO节点暂停后追赶、两个证据端点相关暂停、材料存储不可用及审计持久化故障各执行5次，并在恢复后运行主端点探针、保守终态检查或同回执恢复写入检查。每组运行保存manifest、逐查询JSONL、汇总JSON、CSV及SQLite审计库；独立Python审计器直接读取这些规范化数据库。第二机器复现脚本已固定上述参数并记录环境，但截至本稿尚无第二主机实测结果。",
    )

    # Evaluation overview and setup.
    set_text(
        find_paragraph(doc, "实验围绕五个研究问题展开。"),
        "实验围绕五个研究问题展开。RQ1检验证据规范化、任务绑定、两类真实异构取证和第二业务负载；RQ2评估DAG调度、形态与规模、真实双链端到端延迟、单/多钱包并发饱和、资源利用和专用锚定开销；RQ3评估受控故障、节点暂停、网络分区、相关故障、恢复追赶及材料与审计存储失败下的恢复与有界终结；RQ4以独立实现检验回执完整性、叶级绑定、篡改检测与重构开销；RQ5通过缓存、重试、调度剪枝和服务选择消融量化已实现机制。",
    )
    set_text(
        find_paragraph(doc, "评估使用Fabric与FISCO-BCOS双链环境"),
        f"实验分为真实双链执行、本地控制实验和探索性仿真三层。各层均运行在{meta['cpu']}、{meta['logicalCpus']}逻辑核、{meta['memoryGiB']:.0f} GiB内存的单机上，系统为{meta['platform']}，Node.js版本为{meta['node']}。双链层使用Fabric测试网络、FISCO-BCOS四节点网络、SQLite审计库和专用锚定合约；本地层不访问真实链或外部模型，用于隔离证据验证、DAG形态、并发、审计和机制开销；服务调度可用率来自固定种子的故障模型，只作为探索性结果。",
    )
    set_text(
        find_paragraph(doc, "比较基线包括："),
        "第一双链工作负载为供应链融资10节点DAG，包含正常、Fabric/FISCO候补端点、语义服务超时、Fabric端点耗尽和持续语义分歧六个场景。第二工作负载为监管披露10节点DAG，覆盖正常、超限、主体不一致、披露缺失和端点重试；先以500次确定性本地夹具验证结构复用，再以150次计量查询从Fabric和FISCO读取全部五类证据。DAG扩展实验覆盖宽、平衡和深链三种形态，16/32/64节点及256 B/4 KiB/16 KiB证据；真实双链并发为1/2/4/8/16。",
    )
    set_text(
        find_paragraph(doc, "性能指标包括端到端延迟及"),
        "主要指标为终态正确性、端到端及p50/p95/p99延迟、吞吐、DAG执行与锚定持久化开销、重试次数、回执大小、叶级绑定覆盖和审计核验率。本地调度实验报告5次重复，其受控服务时间只用于比较策略相对差异，不代表真实RPC或模型推理时延。进程RSS差分受垃圾回收影响出现负值，因而不作为论文结果。",
    )
    setup_anchor = find_paragraph(doc, "主要指标为终态正确性")
    insert_caption_and_table(doc, setup_anchor, "表2  实验环境、规模与证据类型", ["项目", "设置"], [
        ["执行环境", f"{meta['cpu']}；{meta['logicalCpus']}逻辑核；{meta['memoryGiB']:.0f} GiB；{meta['node']}；单机回环网络"],
        ["真实双链功能", "供应链6场景×100查询；3次预热；Fabric/FISCO读取、SQLite与FISCO锚定实际执行"],
        ["真实双链并发", "单钱包、四钱包各并发1/2/4/8/16；每档3×20计量+2预热；1 s容器资源采样"],
        ["RQ1证据验证", "200正常样本；8类故障×100；SHA-256与Ed25519实际执行"],
        ["RQ2调度", "8/16/32节点；3种策略；5次运行×10查询；受控本地服务时间"],
        ["RQ3恢复", "8类故障×2策略×100；状态机与恢复分支实际执行"],
        ["RQ4审计", "8/16/32/64节点×100；每份4种篡改；签名、摘要与逆向遍历实际执行"],
        ["物理故障", "Fabric Org1 peer与FISCO node0各暂停20次；无调度器故障注入"],
        ["韧性故障", "网络分区、节点追赶、相关双节点暂停、材料/审计存储失败各5次；恢复探针"],
        ["第二业务负载", "监管披露5场景×100本地控制；5场景×30真实链计量；预设终态"],
        ["扩展性", "3种DAG形态×3规模×3证据体积×100；5档本地并发×200"],
        ["模型基准准备", f"{model_cases['caseCount']}条模型独立规则标签案例；6类规则；人类双标模板已生成；真实推理尚未执行"],
        ["独立审计", f"Python交叉实现复核643份历史v1、{live_v2_queries}份真实链v2与900份本地v2回执；每份4种篡改"],
        ["链上数据", f"正式数据集专用回执锚定{total_formal_anchors}笔；旧路径归档扫描{chain['scannedBlocks']}个区块"],
    ])

    # RQ1.
    set_text(
        find_paragraph(doc, "首先对每类查询逐项检查"),
        "RQ1以统一证据对象为测试单元。正常样本按固定查询、源链、位置和模式生成并签名；异常样本分别修改链标识、区块位置、签名、模式、有效期、必需字段、qid或负载。验证器必须在任何语义节点执行前拒绝异常对象，并保留细分错误码。",
    )
    set_text(
        find_paragraph(doc, "现有原型记录已经证明"),
        f"200个正常证据全部通过，正常接受率为100%（Wilson 95%区间{pct(rq1['validAcceptance']['ci95Lower'], 2)}–100%）；8类故障共800例全部被拒绝，检出率为100%（{pct(rq1['injectedFaultDetection']['ci95Lower'], 2)}–100%）。监管披露的本地控制工作负载500/500均得到预设终态或业务真值。真实链复核进一步将实体、许可和披露存入Fabric，将交易和政策存入FISCO：150/150个计量查询完成且审计通过，正常与端点重试各30次均为SATISFIED，超限与主体不一致各30次均为UNSATISFIED，披露字段缺失30次均保守终结为INSUFFICIENT_EVIDENCE。真实链结果证明第二DAG已接入异构账本，但预置记录及确定性语义服务仍不代表真实监管业务的外部有效性。",
    )
    rq1_anchor = find_paragraph(doc, "200个正常证据全部通过")
    insert_caption_and_table(doc, rq1_anchor, "表3  统一证据验证结果", ["类别", "样本数", "正确处理", "结果"], [
        ["正常证据", 200, 200, "全部接受"],
        ["任务/来源绑定（错链、错位置、错qid）", 300, 300, "全部拒绝"],
        ["完整性（错签名、负载替换）", 200, 200, "全部拒绝"],
        ["模式/时效/完整字段", 300, 300, "全部拒绝"],
        ["监管披露：正常/超限/主体不一致", 300, 300, "符合预设真值"],
        ["监管披露：缺失/端点重试", 200, 200, "保守终结/恢复"],
        ["监管披露真实链计量", 150, 150, "五场景均符合预设终态"],
    ])

    # RQ2.
    set_text(
        find_paragraph(doc, "延迟实验分别改变DAG宽度与深度"),
        "RQ2首先执行离线调度微基准。顺序基线按拓扑逐节点等待；无依赖全并行基线立即启动全部节点；依赖感知调度只在前驱完成后启动就绪节点。三种模式运行相同节点函数与受控服务时间，因而差异主要来自依赖等待和并行策略。",
    )
    lookup = {(row["nodeCount"], row["mode"]): row for row in rq2_rows}
    r32_seq = lookup[(32, "sequential")]["allQueriesMs"]["p95"]
    r32_dag = lookup[(32, "dependency_aware")]["allQueriesMs"]["p95"]
    reduction = (r32_seq - r32_dag) / r32_seq
    set_text(
        find_paragraph(doc, "已有阶段估计中的8秒正常路径"),
        f"依赖感知调度在8、16和32节点下的p95分别为{lookup[(8, 'dependency_aware')]['allQueriesMs']['p95']:.2f}、{lookup[(16, 'dependency_aware')]['allQueriesMs']['p95']:.2f}和{r32_dag:.2f} ms；32节点时相对顺序执行的{r32_seq:.2f} ms降低{pct(reduction)}。进一步的2,700次形态实验表明，在4 KiB证据下，64节点宽、平衡和深链DAG的p95分别为{next(row['p95'] for row in extended['dagShape'] if row['shape']=='wide' and row['nodeCount']==64 and row['evidenceBytes']==4096):.2f}、{next(row['p95'] for row in extended['dagShape'] if row['shape']=='balanced' and row['nodeCount']==64 and row['evidenceBytes']==4096):.2f}和{next(row['p95'] for row in extended['dagShape'] if row['shape']=='deep' and row['nodeCount']==64 and row['evidenceBytes']==4096):.2f} ms，说明关键路径深度比同档证据体积更直接地决定等待时间。无依赖全并行的p95最低，但它会在必要前驱失败时启动不再需要的节点，因此不能单独用正常路径时延判断优劣。",
    )
    set_text(
        find_paragraph(doc, "已有FISCO账本扫描覆盖"),
        f"第一轮真实双链功能基准共完成{live['totalQueries']}个计量查询。正常场景的平均延迟为{live_by['normal']['totalLatencyMs']['mean']:.3f} ms，p50、p95和p99分别为{live_by['normal']['totalLatencyMs']['p50']:.3f}、{live_by['normal']['totalLatencyMs']['p95']:.3f}和{live_by['normal']['totalLatencyMs']['p99']:.3f} ms。其中DAG执行均值为{live_by['normal']['executionLatencyMs']['mean']:.3f} ms，锚定与持久化均值为{live_by['normal']['anchorAndPersistenceMs']['mean']:.3f} ms，占端到端时间约90.5%。单钱包真实饱和基准完成300个计量查询：并发1/2/4/8/16的吞吐分别为{saturation_throughputs} query/s，p95由{saturation['rows'][0]['p95Ms']:.3f}增至{saturation['rows'][-1]['p95Ms']:.3f} ms，并发2首次满足“吞吐增益低于10%且p95增加超过25%”的饱和判据。四钱包轮询对照同样完成300个计量查询，吞吐依次为{multiwallet_throughputs} query/s，饱和点移至并发8。并发4时吞吐由{multiwallet_by_concurrency[4]['singleThroughput']:.3f}提高至{multiwallet_by_concurrency[4]['multiThroughput']:.3f} query/s（提高{pct(multiwallet_by_concurrency[4]['throughputGainRate'])}），p95由{multiwallet_by_concurrency[4]['singleP95Ms']:.3f}降至{multiwallet_by_concurrency[4]['multiP95Ms']:.3f} ms。该对照定位并缓解了单账户nonce串行瓶颈；额外钱包为实验临时签名者，两组运行日期不同且未随机交错，不能外推为平台上限。FISCO容器平均CPU为{min(row['resources']['containers']['qkl-fisco-nodes']['cpuPercent']['mean'] for row in saturation['rows']):.1f}%–{max(row['resources']['containers']['qkl-fisco-nodes']['cpuPercent']['mean'] for row in saturation['rows']):.1f}%；主机同时存在外部负载，故不使用host load作为归因证据。",
    )
    set_text(
        find_paragraph(doc, "账本样本的区块间隔中位数约为"),
        f"正式真实链审计数据库共保存{total_formal_anchors}笔锚定交易，其中历史功能与物理故障643笔、单钱包饱和实验310笔、监管披露151笔、韧性与材料存储故障40笔、四钱包对照{multiwallet_anchors}笔；所有记录的gasUsed均为{live_anchor['gasMean']:,.0f}。单钱包饱和实验的310个唯一交易哈希连续覆盖FISCO区块690–999，监管披露的151个唯一交易哈希连续覆盖区块{regulatory_anchor['heightMin']}–{regulatory_anchor['heightMax']}，韧性与材料存储故障的40个唯一交易哈希覆盖区块{resilience_anchor['heightMin']}–{material_outage_anchor['heightMax']}，四钱包对照的{multiwallet_anchors}个唯一交易哈希覆盖区块{min(item['heightMin'] for item in multiwallet_anchor_stats)}–{max(item['heightMax'] for item in multiwallet_anchor_stats)}。合约每次接收qid摘要、queryDigest、receiptRoot与两个uint8终态字段，因此单次锚定不随证据数或语义服务数逐项增长。旧路径的receiveLite、submitBlockHeader和send仅作历史参考，不与专用合约数据混合。",
    )
    rq2_anchor = find_paragraph(doc, "依赖感知调度在8、16和32节点")
    scheduler_table = []
    for node_count in [8, 16, 32]:
        seq = lookup[(node_count, "sequential")]["allQueriesMs"]["p95"]
        eager = lookup[(node_count, "eager_no_dependency")]["allQueriesMs"]["p95"]
        dag = lookup[(node_count, "dependency_aware")]["allQueriesMs"]["p95"]
        scheduler_table.append([node_count, f"{seq:.2f}", f"{eager:.2f}", f"{dag:.2f}", pct((seq - dag) / seq)])
    insert_caption_and_table(doc, rq2_anchor, "表4  离线DAG调度微基准（p95，ms）", ["节点数", "顺序", "无依赖全并行", "依赖感知", "相对顺序降幅"], scheduler_table)
    live_anchor_paragraph = find_paragraph(doc, "第一轮真实双链功能基准共完成")
    insert_caption_and_table(doc, live_anchor_paragraph, "表6  单钱包与四钱包真实双链饱和对照", ["并发", "单钱包吞吐", "四钱包吞吐", "单钱包p95(ms)", "四钱包p95(ms)", "四钱包排队p95(ms)"], [
        [single["concurrency"], f"{single['throughput']:.3f}", f"{multi['throughput']:.3f}", f"{single['p95Ms']:.3f}", f"{multi['p95Ms']:.3f}", f"{multi['anchorQueueP95Ms']:.3f}"]
        for single, multi in zip(saturation["rows"], multiwallet["rows"])
    ])
    insert_caption_and_table(doc, live_anchor_paragraph, "表5  真实双链端到端结果", ["场景", "n", "终态", "p50(ms)", "p95(ms)", "平均重试", "审计通过"], [
        ["正常", 100, "ANCHORED", live_by["normal"]["totalLatencyMs"]["p50"], live_by["normal"]["totalLatencyMs"]["p95"], 0, "100/100"],
        ["Fabric候补端点", 100, "ANCHORED", live_by["fabric_endpoint_retry"]["totalLatencyMs"]["p50"], live_by["fabric_endpoint_retry"]["totalLatencyMs"]["p95"], 1, "100/100"],
        ["FISCO候补端点", 100, "ANCHORED", live_by["fisco_endpoint_retry"]["totalLatencyMs"]["p50"], live_by["fisco_endpoint_retry"]["totalLatencyMs"]["p95"], 1, "100/100"],
        ["语义服务超时", 100, "ANCHORED", live_by["semantic_service_timeout"]["totalLatencyMs"]["p50"], live_by["semantic_service_timeout"]["totalLatencyMs"]["p95"], 1, "100/100"],
        ["Fabric候补耗尽", 100, "FAILED", live_by["fabric_endpoint_exhausted"]["totalLatencyMs"]["p50"], live_by["fabric_endpoint_exhausted"]["totalLatencyMs"]["p95"], 1, "100/100"],
        ["语义持续分歧", 100, "INSUFFICIENT", live_by["semantic_disagreement"]["totalLatencyMs"]["p50"], live_by["semantic_disagreement"]["totalLatencyMs"]["p95"], 1, "100/100"],
    ])

    # RQ3.
    set_text(
        find_paragraph(doc, "基础设施故障通过关闭源链节点"),
        "RQ3分为状态机注入、真实双链调度边界注入和基础设施故障三层。前两层覆盖单/双端点不可用、索引滞后、证据过期或缺失、输入根错误、服务超时和语义分歧；基础设施层通过Docker暂停节点、断开容器网络和恢复连接，并在锚定成功后注入一次审计持久化失败。",
    )
    no_recovery = rq3["aggregate"]["noRecoveryCorrectTerminal"]
    bounded = rq3["aggregate"]["boundedRecoveryCorrectTerminal"]
    set_text(
        find_paragraph(doc, "证据故障包括过期时间戳"),
        f"离线对照中，无恢复策略仅有300/800例进入预期终态（{pct(no_recovery['rate'])}）；有界策略的800/800例均进入预期终态。真实双链受控场景中，Fabric与FISCO单端点失败各100次均通过第二候选恢复，语义服务超时100次均切换至第二服务；Fabric端点耗尽的100个查询全部进入FAILED，持续语义分歧的100个查询全部进入INSUFFICIENT_EVIDENCE。",
    )
    set_text(
        find_paragraph(doc, "语义服务故障通过增加响应延迟"),
        f"容器级暂停实验中，Fabric Org1 peer的20个查询全部切换至Org2 peer，p95为{physical_by['fabric_peer_paused']['p95Ms']:.3f} ms；FISCO node0的20个查询全部切换至node1，p95为{physical_by['fisco_node_paused']['p95Ms']:.3f} ms。新增韧性实验的五个场景各执行5次：Fabric网络分区均切换并在重连探针中恢复主peer，p95为{resilience_by['fabric_network_partition']['p95Ms']:.3f} ms；FISCO node0暂停均切换，恢复追赶p95为{resilience_by['fisco_node_resync']['resyncP95Ms']:.3f} ms，且探针重新使用node0；同时暂停两个已配置FISCO证据端点时5/5查询保守进入FAILED，恢复后原回执被锚定且探针通过，p95为{resilience_by['fisco_correlated_pause']['p95Ms']:.3f} ms；材料存储不可用时5/5查询保守终结并锚定可审计回执，p95为{material_outage_result['p95Ms']:.3f} ms；审计存储fail-once场景5/5使用同一回执和同一链上交易恢复写入，没有重复锚定，p95为{resilience_by['audit_store_outage']['p95Ms']:.3f} ms。材料存储故障是针对单一文档候选的调度边界注入，原型尚未部署独立材料服务。FISCO周期性约6 s长尾已由日志归因为暂停提议者触发3,000 ms PBFT超时和视图切换，叠加约1.5 s失效证据端点超时；该归因只适用于当前四节点单机配置。",
    )
    rq3_anchor = find_paragraph(doc, "离线对照中")
    insert_caption_and_table(doc, rq3_anchor, "表7  故障恢复与保守终结", ["实验层", "场景", "正确行为", "结果"], [
        ["离线", "8类故障×有界策略", "800/800", "预期终态"],
        ["双链受控", "三类可恢复故障", "300/300", "恢复后锚定"],
        ["双链受控", "端点耗尽/持续分歧", "200/200", "FAILED/INSUFFICIENT"],
        ["容器级", "Fabric/FISCO单节点暂停", "40/40", "候补端点切换"],
        ["容器级", "Fabric网络分区/FISCO恢复追赶", "10/10", "切换、追赶和主端点探针"],
        ["容器级", "两个FISCO证据端点相关暂停", "5/5", "FAILED后恢复锚定"],
        ["调度边界", "材料存储不可用", "5/5", "保守终结并锚定"],
        ["持久化", "锚定后审计存储fail-once", "5/5", "同回执恢复、无重复交易"],
    ])

    # RQ4.
    set_text(
        find_paragraph(doc, "审计实验从已完成查询中随机抽样"),
        "RQ4包含两类验证。首先，对8、16、32和64节点DAG各生成100份签名回执，并用原审计器核验400份完整回执及1,600个篡改实例。其次，使用不复用Node.js审计代码的Python实现，从规范化SQLite表独立重构真实双链、物理故障、并发、监管和韧性回执，重新计算节点、依赖、证据、服务输出和回执根，并分别注入删除节点、替换输出、错误版本和删除证据四类篡改。该交叉实现用于降低生成器与校验器共享实现导致的同源偏差。",
    )
    audit_agg = rq4["aggregate"]
    largest = next(row for row in rq4["rows"] if row["nodeCount"] == 64)
    set_text(
        find_paragraph(doc, "重构开销实验改变DAG和证据包规模"),
        f"同源实验的400/400份完整回执均成功重构（Wilson 95%区间{pct(audit_agg['cleanReconstruction']['ci95Lower'], 2)}–100%），1,600/1,600个篡改均被检测并定位。独立实现对真实双链603份和物理故障40份历史回执全部重构并与643笔锚点匹配，2,572个篡改实例全部检出并定位；两库重构p95分别为{historical_live['reconstructionMs']['p95']:.3f}和{historical_physical['reconstructionMs']['p95']:.3f} ms。交叉核验同时暴露历史格式局限：成功证据节点可通过负载摘要推断绑定的覆盖率为{pct(historical_binding_rate)}，显式服务节点绑定率为0%。修复后的v2格式首先在900份本地控制回执上实现4,300个成功证据节点和800个服务输出100%显式绑定；随后对单钱包饱和、监管、韧性和四钱包对照共{live_v2_queries}份真实链回执进行独立复核，全部完成重构和锚点匹配，{live_v2_evidence:,}个成功证据节点和{live_v2_services}个服务输出显式绑定率均为100%，{live_v2_mutations:,}个篡改全部检出并定位。四钱包的五个数据库各有62份回执，310/310份均通过重构和锚点核对；历史锚点仅能支持整体完整性结论，不能追溯声称已具备100%叶级显式归因。",
    )
    rq4_anchor = find_paragraph(doc, "同源实验的400/400份完整回执均成功重构")
    insert_caption_and_table(doc, rq4_anchor, "表8  回执审计、锚点核对与叶级绑定", ["数据集", "回执", "重构/根匹配", "叶级绑定", "篡改检测/定位"], [
        ["生成回执（同源审计器）", 400, "400/400；无链锚点", "节点/签名检查", "1,600/1,600"],
        ["历史双链+物理故障（独立审计器）", 643, "643/643；643/643", f"证据{pct(historical_binding_rate)}；服务显式0%", "2,572/2,572"],
        ["v2单/多钱包并发+监管+韧性（独立审计器）", live_v2_queries, f"{live_v2_queries}/{live_v2_queries}；{live_v2_queries}/{live_v2_queries}", "证据/服务显式100%", f"{live_v2_mutations:,}/{live_v2_mutations:,}"],
        ["v2本地控制回执（独立审计器）", 900, "900/900；900/900本地锚点", "证据/服务显式100%", "3,600/3,600"],
    ])

    # RQ5.
    set_text(
        find_paragraph(doc, "DAG调度消融比较顺序执行"),
        f"在共享真实性门控失败时，依赖感知调度分别在8、16和32节点DAG中避免3、5和9次后继执行，占全部节点的{pct(pruning[0]['avoidedRate'])}、{pct(pruning[1]['avoidedRate'])}和{pct(pruning[2]['avoidedRate'])}；无依赖全并行仍会启动全部节点。结合表4可见，全并行以额外工作换取最低正常路径时延，依赖感知策略则在相对顺序执行显著降时延的同时保留失败剪枝。",
    )
    set_text(
        find_paragraph(doc, "可靠性机制消融依次关闭缓存"),
        f"缓存消融中，关闭缓存的100次查询产生500次远端夹具读取，p95为{extended['cacheAblation'][0]['p95']:.3f} ms；启用版本化缓存后仅首次查询产生5次读取，其余命中495次，p95降至{extended['cacheAblation'][1]['p95']:.3f} ms。重试消融对同一可恢复端点故障各执行100次：关闭恢复时成功率为{pct(extended['retryAblation'][0]['querySuccessRate'])}，有界候补端点启用后为{pct(extended['retryAblation'][1]['querySuccessRate'])}，代价是平均节点尝试数由{extended['retryAblation'][0]['meanNodeAttempts']:.0f}增至{extended['retryAblation'][1]['meanNodeAttempts']:.0f}。这些结果来自本地确定性夹具，量化机制方向而非真实网络绝对收益。确定性门控消融由RQ1给出：错误证据均在语义调用前被拒绝；本轮未接入真实LLM，故不报告token节省或模型误放行率。",
    )
    set_text(
        find_paragraph(doc, "服务治理消融比较固定服务"),
        f"探索性服务调度在共享的1,000组种子故障样本上比较固定单服务、可靠性单服务和可靠性—负载—多样性双服务。前两者均持续选择同一提供者，可用率为{pct(extended['serviceSchedulingAblation'][0]['availabilityRate'])}、最大提供者份额100%；组合策略将可用率提高至{pct(extended['serviceSchedulingAblation'][2]['availabilityRate'])}，最大份额降至{pct(extended['serviceSchedulingAblation'][2]['maxProviderShare'])}，HHI由1.000降至{extended['serviceSchedulingAblation'][2]['hhi']:.3f}。该结果来自预设可靠性与模型族相关故障，并非真实LLM测量。既有16个困难语义案例、256次历史真实模型推理仅支持模型族多样性的设计动机。本文另构建{model_cases['caseCount']}条模型输出无关的确定性规则标签案例，覆盖6类业务规则并提供人类双标模板，运行器可报告校准、误放、误拒和跨模型错误相关性；但当前机器无可用模型后端，人类双标也未完成，故不报告该扩展基准的模型准确率。",
    )
    rq5_anchor = find_paragraph(doc, "在共享真实性门控失败时")
    insert_caption_and_table(doc, rq5_anchor, "表9  共享门控失败时的DAG剪枝消融", ["节点数", "全并行执行数", "依赖感知执行数", "避免执行", "避免比例"], [
        [row["nodeCount"], row["eagerExecuted"], row["dependencyAwareExecuted"], row["avoidedExecutions"], pct(row["avoidedRate"])] for row in pruning
    ])

    # Evidence boundary, validity, and conclusion.
    set_text(find_paragraph(doc, "7.7 当前证据与完整结论的界限"), "7.7 实验证据分层与结论边界")
    set_text(
        find_paragraph(doc, "当前材料可以支持三项有限结论"),
        "证据分层后可以支持以下结论：两类业务的真实双链取证、DAG执行、SQLite持久化和专用回执锚定已在同一路径中连通；供应链600个受控场景和监管披露150个真实链计量查询均得到预期终态；40次单物理节点暂停均完成候补切换。25次韧性场景验证了所测Fabric网络分区、FISCO恢复追赶、两个已配置证据端点相关暂停、材料存储不可用和锚定后审计持久化失败的恢复或保守终结行为。单钱包和四钱包各300次真实并发查询表明，轮询四个独立账户可将饱和点由并发2移至并发8，并在并发4将吞吐提高至7.150 query/s。正式数据集1,454笔专用锚定交易给出定长摘要写入成本；独立实现完成643份历史v1和811份真实链v2回执的根级交叉核验，并确认v2叶级显式绑定为100%。本地控制实验支持DAG形态、缓存、重试和机制消融；种子故障模型只支持服务调度的探索性比较。120条扩展模型案例及运行器已就绪，但无扩展真实推理结果。历史回执的叶级绑定缺口已经披露，不能用v2结果追溯替代。尚不能据此声称跨地域容量、一般多节点或拜占庭容错、外部材料长期可用性或语义真值正确率。",
    )
    set_text(
        find_paragraph(doc, "性能结果容易受到区块配置"),
        "性能结果受区块配置、网络位置、模型负载、Node.js定时器和缓存预热影响。真实双链实验位于单机回环网络；单钱包基准识别当前nonce串行瓶颈，四钱包对照使用同一主机但运行日期不同、未随机交错，临时签名账户也不是生产密钥管理方案。第二机器复现脚本已准备，但尚无独立主机结果，故不能声称跨机器外部有效性。资源采样期间主机存在外部负载，故仅报告容器级CPU/内存，不以host load归因。监管场景虽读取真实账本状态，但记录由实验预置，语义服务仍为确定性本地实现。相关故障仅暂停两个已配置FISCO证据端点，四节点网络和锚定端点仍存活；材料存储故障是单一文档候选上的确定性调度边界注入，原型尚未部署独立材料服务；审计存储故障采用锚定后的确定性fail-once注入，不等同于物理磁盘、文件系统损坏或长期材料不可用。FISCO长尾归因来自当前四节点PBFT配置，不能外推为其他部署的固定延迟。120条模型基准案例由显式规则自动标注，尚无人类双标和扩展真实模型推理，不能据此报告校准或一般业务错误率。独立Python审计器降低了同源实现偏差，但仍复用了论文定义的规范序列化规则，四类篡改不能代表未知攻击全集。历史回执仅有79.4%的成功证据节点可推断绑定且无服务输出显式nodeId；v2修复不改变历史记录。Fabric当前证据定位也不构成交易包含证明。审计完整性不等同于材料可用性或业务结论正确性。",
    )
    set_text(
        find_paragraph(doc, "本文研究异构联盟链上的跨链业务语义查询"),
        "本文研究异构联盟链上的跨链业务语义查询：调用方合约需要组合多源账本状态和链下材料判断业务谓词，而现有跨链传输和单次语义服务调用缺少统一证据语境、失败处置和可重构审计。为此，本文提出统一证据对象和证据依赖图，以异构适配、并行取证、确定性预验证、受证据约束的语义服务、冲突升级和摘要锚定构成端到端框架。",
    )
    set_text(
        find_paragraph(doc, "分析表明，在明确的密码学"),
        "在明确的密码学、终局性和数据可用性假设下，框架能够检测跨任务或跨版本证据替换，并在有限预算内终结。两类真实双链负载验证了从异构取证、DAG执行到审计持久化和摘要锚定的路径；节点暂停、网络分区、恢复追赶、所测相关端点故障、材料存储不可用和持久化fail-once实验验证了候补切换、保守终结与无重复锚定恢复。单/四钱包对照定位了单账户nonce串行瓶颈，并验证多签名账户轮询可在本原型上提高吞吐和推迟饱和。独立审计确认历史回执根和锚点一致并识别出叶级绑定不足，也确认v2真实链回执已补齐显式节点绑定。上述证据仍未覆盖跨地域部署、一般多节点或拜占庭故障、材料长期可用或扩展真实模型的语义真值保证。框架的贡献是为智能合约消费链下语义输出提供可验证的证据边界、保守终态和事后责任材料，而不是把服务一致性等同于业务真实性。",
    )

    # Core architecture figures use the repo-native redrawn assets, not the
    # legacy committee figures embedded in the earlier manuscript.
    insert_picture(
        doc,
        find_paragraph(doc, "本节给出框架从查询登记到结果交付的完整路径"),
        FIGURE_DIR / "1_figure_auditable_semantic_query_main_architecture.png",
        "图1  面向异构联盟链的可审计语义查询总体架构",
    )
    insert_picture(
        doc,
        find_paragraph(doc, "编译器对DAG进行模式检查和闭环检查"),
        FIGURE_DIR / "2_figure_supply_chain_financing_evidence_dependency.png",
        "图3  供应链融资查询的证据依赖关系",
    )
    insert_picture(
        doc,
        find_paragraph(doc, "当所有必需分支得到可用结果时"),
        FIGURE_DIR / "4_figure_dag_state_machine_combined.png",
        "图2  DAG执行与有界查询状态机",
    )
    insert_picture(
        doc,
        find_paragraph(doc, "正式真实链审计数据库共保存"),
        FIGURE_DIR / "5_figure_performance_scalability.png",
        "图4  真实双链功能、物理故障与并发饱和结果",
    )
    insert_picture(
        doc,
        find_paragraph(doc, "同源实验的400/400份完整回执均成功重构"),
        FIGURE_DIR / "7_figure_independent_audit_binding.png",
        "图5  独立回执核验与历史v1/真实链v2叶级绑定覆盖",
    )
    insert_picture(
        doc,
        find_paragraph(doc, "缓存消融中"),
        FIGURE_DIR / "6_figure_dag_mechanism_ablation.png",
        "图6  本地DAG扩展与机制消融（控制实验/探索性仿真）",
    )

    # Normalize the bibliography to the numbered MDPI/ACS pattern and retain
    # DOI or accessed-date locators verified against publisher/standards pages.
    references = [
        "[1] Belchior, R.; Vasconcelos, A.; Guerreiro, S.; Correia, M. A Survey on Blockchain Interoperability: Past, Present, and Future Trends. ACM Comput. Surv. 2021, 54, 168. https://doi.org/10.1145/3471140.",
        "[2] Zamyatin, A.; Harz, D.; Lind, J.; Panayiotou, P.; Gervais, A.; Knottenbelt, W. XCLAIM: Trustless, Interoperable, Cryptocurrency-Backed Assets. In Proceedings of the 2019 IEEE Symposium on Security and Privacy, San Francisco, CA, USA, 19–23 May 2019; pp. 193–210. https://doi.org/10.1109/SP.2019.00085.",
        "[3] Liu, Z.; Xiang, Y.; Shi, J.; Gao, P.; Wang, H.; Xiao, X.; Wen, B.; Hu, Y.-C. HyperService: Interoperability and Programmability across Heterogeneous Blockchains. In Proceedings of the 2019 ACM SIGSAC Conference on Computer and Communications Security, London, UK, 11–15 November 2019; pp. 549–566. https://doi.org/10.1145/3319535.3355503.",
        "[4] Hyperledger Foundation. Hyperledger Cacti: A Pluggable Interoperability Framework for Blockchain and DLT Networks. Available online: https://hyperledger-cacti.github.io/cacti/ (accessed on 1 September 2026).",
        "[5] Hyperledger Foundation. Levels of Interoperability. Available online: https://hyperledger-cacti.github.io/cacti/weaver/what-is-interoperability/levels-of-interoperability/ (accessed on 1 September 2026).",
        "[6] Caldarelli, G. Understanding the Blockchain Oracle Problem: A Call for Action. Information 2020, 11, 509. https://doi.org/10.3390/info11110509.",
        "[7] Mühlberger, R.; Bachhofner, S.; Castelló Ferrer, E.; Di Ciccio, C.; Weber, I.; Wöhrer, M.; Zdun, U. Foundational Oracle Patterns: Connecting Blockchain to the Off-Chain World. In Business Process Management: Blockchain and Robotic Process Automation Forum; Springer: Cham, Switzerland, 2020; Volume 393, pp. 35–51. https://doi.org/10.1007/978-3-030-58779-6_3.",
        "[8] Caldarelli, G. Can Artificial Intelligence Solve the Blockchain Oracle Problem? Unpacking the Challenges and Possibilities. Front. Blockchain 2025, 8, 1682623. https://doi.org/10.3389/fbloc.2025.1682623.",
        "[9] Mendling, J.; Weber, I.; van der Aalst, W.M.P.; vom Brocke, J.; Cabanillas, C.; Daniel, F.; Debois, S.; Di Ciccio, C.; Dumas, M.; Dustdar, S.; Gal, A.; García-Bañuelos, L.; Governatori, G.; Hull, R.; La Rosa, M.; Leopold, H.; Leymann, F.; Recker, J.; Reichert, M.; Reijers, H.A.; Rinderle-Ma, S.; Solti, A.; Rosemann, M.; Schulte, S.; Singh, M.P.; Slaats, T.; Staples, M.; Weber, B.; Weidlich, M.; Weske, M.; Xu, X.; Zhu, L. Blockchains for Business Process Management—Challenges and Opportunities. ACM Trans. Manag. Inf. Syst. 2018, 9, 4. https://doi.org/10.1145/3183367.",
        "[10] Di Ciccio, C.; Meroni, G.; Plebani, P. On the Adoption of Blockchain for Business Process Monitoring. Softw. Syst. Model. 2022, 21, 915–937. https://doi.org/10.1007/s10270-021-00959-x.",
        "[11] Moreau, L.; Missier, P., Eds. PROV-DM: The PROV Data Model; W3C Recommendation: 30 April 2013. Available online: https://www.w3.org/TR/2013/REC-prov-dm-20130430/ (accessed on 1 September 2026).",
        "[12] Liang, X.; Shetty, S.; Tosh, D.K.; Kamhoua, C.A.; Kwiat, K.A.; Njilla, L. ProvChain: A Blockchain-Based Data Provenance Architecture in Cloud Environment with Enhanced Privacy and Availability. In Proceedings of the 17th IEEE/ACM International Symposium on Cluster, Cloud and Grid Computing, Madrid, Spain, 14–17 May 2017; pp. 468–477. https://doi.org/10.1109/CCGRID.2017.8.",
        "[13] Pan, B.; Stakhanova, N.; Ray, S. Data Provenance in Security and Privacy. ACM Comput. Surv. 2023, 55, 323. https://doi.org/10.1145/3593294.",
        "[14] Hyperledger Fabric. Transaction Flow, Release 2.2. Available online: https://hyperledger-fabric.readthedocs.io/en/release-2.2/txflow.html (accessed on 1 September 2026).",
    ]
    for index, reference in enumerate(references, start=1):
        set_text(find_paragraph(doc, f"[{index}]"), reference)

    # Normalize spacing in newly rewritten body text.
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.25

    doc.core_properties.title = "面向异构联盟链的可审计跨链语义查询框架（统一实证修订版）"
    doc.core_properties.subject = "整合真实双链、物理与韧性故障、本地扩展、独立审计与机制消融结果"
    doc.core_properties.comments = "Unified evidence revision; historical receipt binding limitation disclosed; evidence tiers preserved."
    doc.save(OUTPUT)

    changelog = f"""# 修订说明

## 输出

- 修订稿：`{OUTPUT.name}`
- 原始实验结果：`experiments/semantic_query_framework_results.json`
- 可复现实验脚本：`experiments/semantic_query_framework_experiments.js`
- 分项CSV：`semantic_query_dag_results.csv`、`semantic_query_fault_results.csv`、`semantic_query_audit_results.csv`
- 扩展实验：`experiments/semantic-query-extended-evaluation/results.json`
- 独立审计：`experiments/independent_audit_verification.json`、`experiments/semantic-query-receipt-v2-independent-verification.json`

## 主要修改

1. 删除{removed_notes}个“作者注释/实验证据缺口”单元格，保留统一证据对象字段表。
2. 将摘要、贡献、原型实现、实验、局限和结论从“待补实验”改为有证据边界的完成时表述。
3. 合入600个真实双链查询，报告终态、p50/p95/p99、重试和审计核验。
4. 合入40次容器级Fabric/FISCO节点暂停实验，并将周期性长尾归因到PBFT提议者超时和视图切换。
5. 使用1,454笔SemanticQueryAuditAnchor交易替代旧receiveLite小样本作为当前锚定成本证据。
6. 增加监管披露第二业务负载500次、DAG形态/规模/证据体积2,700次及本地并发1,000次实验。
7. 增加缓存、重试和服务调度消融，并将服务调度明确标注为种子故障模型上的探索性结果。
8. 使用独立Python实现交叉核验643份历史回执，披露历史证据绑定79.4%、服务显式绑定0%的格式缺口。
9. 验证修复后的v2格式：900份本地回执、4,300个证据节点和800个服务输出均为100%显式绑定。
10. 插入新的总架构、供应链证据依赖、DAG状态机及3张结果图，不复用旧Agent委员会图。
11. 增加并发1/2/4/8/16的真实双链饱和实验300次，采集锚定排队和Docker容器资源，定位单钱包锚定瓶颈。
12. 将监管披露五类证据全部接入Fabric/FISCO状态，完成150次正式计量查询及151份回执独立审计。
13. 独立复核811份真实链v2回执、4,010个成功证据节点和776个服务输出，显式绑定与3,244次篡改检测/定位均为100%。
14. 按Electronics/MDPI编号制统一14条参考文献，并逐条核验题名、作者、卷期页、DOI/URL与访问日期。
15. 增加网络分区、FISCO恢复追赶、相关双节点暂停、材料存储不可用和审计持久化失败各5次；25/25次得到预期恢复或保守终态，40份回执均通过独立审计。
16. 增加四钱包轮询锚定及300次正式并发对照：并发4吞吐提高至7.150 query/s，饱和点由并发2移至并发8；310份回执均通过独立审计。
17. 生成120条模型输出无关的规则标签案例、人类双标模板及校准/误放/误拒/错误相关性运行器；因当前无模型后端，不虚报扩展推理结果。
18. 增加第二机器一键复现脚本、环境清单、独立审计和校验和流程；外部主机结果仍待执行。

## 后续扩展（非P0）

- 在第二机器或跨地域环境复现，并扩展到更多相关故障与拜占庭故障模型。
- 在120条规则标签案例上实际运行至少两个真实模型，并完成人类双标，报告校准、误放/误拒和跨模型错误相关性。
"""
    CHANGELOG.write_text(changelog, encoding="utf-8")
    print(OUTPUT)
    print(CHANGELOG)


if __name__ == "__main__":
    main()
