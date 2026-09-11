#!/usr/bin/env python3
"""Build an NDSS-oriented review DOCX from the Chinese Markdown manuscript."""

from __future__ import annotations

import argparse
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+\*)")
TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")
LIST_ITEM = re.compile(r"^\s*(?:[-+*]|\d+\.)\s+(.+)$")


def set_run_font(run, *, size=10, bold=None, italic=None, code=False):
    run.font.name = "Courier New" if code else "Times New Roman"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def add_inline(paragraph, text, *, size=10):
    pos = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos : match.start()]), size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=max(8, size - 1), code=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, italic=True)
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), size=size)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_columns(section, count=2, space_twips=360):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    node = cols[0] if cols else OxmlElement("w:cols")
    node.set(qn("w:num"), str(count))
    node.set(qn("w:space"), str(space_twips))
    if not cols:
        sect_pr.append(node)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    for name, size in (("Title", 16), ("Heading 1", 12), ("Heading 2", 11), ("Heading 3", 10)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(3)

    if "Code Block" not in [style.name for style in doc.styles]:
        style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Courier New"
        style.font.size = Pt(8)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "等线")
        style.paragraph_format.space_after = Pt(2)


def markdown_cells(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def joined_paragraph(lines):
    result = ""
    for raw in lines:
        item = raw.strip()
        if not result:
            result = item
        elif re.search(r"[A-Za-z0-9`\]\)]$", result) and re.match(r"^[A-Za-z0-9`(\[]", item):
            result += " " + item
        else:
            result += item
    return result


def add_table(doc, lines):
    rows = [markdown_cells(line) for line in lines if not TABLE_SEPARATOR.match(line)]
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(width):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[c_idx] if c_idx < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, text, size=7)
            if r_idx == 0:
                set_cell_shading(cell, "D9E2F3")
                for run in paragraph.runs:
                    run.bold = True
            tc_pr = cell._tc.get_or_add_tcPr()
            cant_split = OxmlElement("w:cantSplit")
            tc_pr.append(cant_split)


def add_code(doc, text):
    paragraph = doc.add_paragraph(style="Code Block")
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text.rstrip())
    set_run_font(run, size=8, code=True)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)


def add_body_paragraph(doc, text, style=None):
    paragraph = doc.add_paragraph(style=style)
    if text.startswith("[") and re.match(r"^\[\d+\]", text):
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(paragraph, text)
    return paragraph


def convert(source: Path, output: Path):
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "面向跨链语义查询的智能合约可信调用链下智能体协议"
    doc.core_properties.subject = "NDSS 双盲评审内容稿（中文版）"
    doc.core_properties.author = "Anonymous"
    doc.core_properties.last_modified_by = "Anonymous"
    doc.core_properties.comments = "Generated from the anonymized Markdown manuscript."
    doc.core_properties.keywords = "cross-chain, multi-agent, Byzantine robustness, auditability"

    lines = source.read_text(encoding="utf-8").splitlines()
    paragraph_buffer = []
    code_buffer = []
    in_code = False
    two_columns = False
    i = 0

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            add_body_paragraph(doc, joined_paragraph(paragraph_buffer))
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                add_code(doc, "\n".join(code_buffer))
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and TABLE_SEPARATOR.match(lines[i + 1]):
            flush_paragraph()
            table_lines = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2)
            if text == "I. 引言" and not two_columns:
                section = doc.add_section(WD_SECTION.CONTINUOUS)
                section.page_width = Inches(8.5)
                section.page_height = Inches(11)
                section.top_margin = Inches(0.72)
                section.bottom_margin = Inches(0.72)
                section.left_margin = Inches(0.68)
                section.right_margin = Inches(0.68)
                set_columns(section, 2)
                two_columns = True
            if level == 1:
                paragraph = doc.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, text, size=16)
            else:
                style = f"Heading {min(level - 1, 3)}"
                doc.add_paragraph(text, style=style)
            i += 1
            continue

        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            flush_paragraph()
            image_path = (source.parent / image.group(2)).resolve()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if image_path.exists():
                paragraph.add_run().add_picture(str(image_path), width=Inches(3.25 if two_columns else 6.7))
            else:
                add_inline(paragraph, f"[缺失图片：{image.group(1)}]")
            i += 1
            continue

        if stripped in {"---", "***", "___"}:
            flush_paragraph()
            i += 1
            continue

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            add_body_paragraph(doc, stripped.lstrip("> ").strip(), style="Quote")
            i += 1
            continue

        item = LIST_ITEM.match(line)
        if item:
            flush_paragraph()
            style = "List Number" if re.match(r"^\s*\d+\.", line) else "List Bullet"
            add_body_paragraph(doc, item.group(1), style=style)
            i += 1
            continue

        paragraph_buffer.append(line)
        i += 1

    flush_paragraph()
    if code_buffer:
        add_code(doc, "\n".join(code_buffer))

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    doc.save(temporary)
    os.replace(temporary, output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    convert(args.source.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
